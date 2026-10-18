import re

from dataclasses import dataclass
from datetime import date, datetime, time
from typing import IO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from app.core.exceptions import InvalidStatementError
from app.imports.parsers.base import ParsedTransaction


HEADER_IDX = 2
DATE_FORMAT = "%Y-%m-%d"

@dataclass
class HyundaiTransaction:
    transferred_at: date
    confirmed_at: date
    holder_type: str
    card_name: str
    merchant: str
    business_no: str
    approval_no: str
    installment_months: int
    amount: int


def digits_only(value)->str:
    return re.sub(r'[^0-9]', '', str(value or ''))

def to_int(value)->int:
    if isinstance(value, int):
        return value
    text = str(value or "").strip()
    digits = digits_only(text)
    if not digits:
        return 0
    number = int(digits)
    return -number if text.startswith("-") else number

def open_excel(file)->list[HyundaiTransaction]:
    """

    :return:
    """
    try:
        wb = load_workbook(file)
        if not wb:
            raise InvalidStatementError(
                "시트를 읽을 수 없습니다. 'Excel 통합 문서(.xlsx)' 형식으로 저장 후 다시 시도하세요."
            )
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        headers = [str(c).replace("\n", " ").strip() for c in rows[HEADER_IDX]]
        card_transaction = []
        for row in rows[HEADER_IDX+1:]:
            if all(c is None for c in row):
                continue
            if str(row[0]).startswith("총 합계"):
                break
            record = dict(zip(headers, row))
            card_transaction.append(to_dataclass(record))
    except (BadZipFile, OSError, KeyError, IndexError, ValueError, TypeError) as e:
        raise InvalidStatementError("현대카드 이용내역 형식이 아닙니다.") from e
    return card_transaction

def parse_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, (int, float)):
        return from_excel(value).date()
    return datetime.strptime(str(value).strip(), DATE_FORMAT).date()

def to_dataclass(record: dict) -> HyundaiTransaction:
    return HyundaiTransaction(
        transferred_at=parse_date(record["이용일"]),
        confirmed_at=parse_date(record["확정일"]),
        holder_type=record["카드구분"],
        card_name=record["카드명(카드번호 뒤 4자리)"],
        merchant=record["가맹점명"],
        business_no=record["사업자번호"],
        approval_no=digits_only(record["승인번호"]),
        installment_months=to_int(record["할부 개월"]),
        amount=to_int(record["이용금액"]),
    )

def to_parsed(t: HyundaiTransaction)->ParsedTransaction:
    return ParsedTransaction(
        transacted_at=t.transferred_at,
        transacted_time=None,
        merchant=t.merchant,
        amount=t.amount,
        discount=0,
        payment_type=to_payment_type(t.installment_months),
        status=to_status(t),
        card_name=t.card_name,
        approval_no=t.approval_no,
    )

def to_payment_type(months:int)->str:
    return "일시불" if months == 0 else f"할부 {months}개월"

def to_status(t: HyundaiTransaction) -> str:
    return "전표매입" if t.amount > 0 else "승인취소"

class HyundaiParser:
    source = "hyundai"

    def parse(self, file:IO[bytes]) -> list[ParsedTransaction]:
        return [to_parsed(t) for t in open_excel(file)]


